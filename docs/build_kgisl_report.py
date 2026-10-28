from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "KGiSL_Project_Report.md"
OUTPUT = ROOT / "Smart_Gate_Access_Control_Project_Report_KGiSL.docx"
TITLE = "SMART AUTOMATED GATE CONTROL AND ACCESS MONITORING SYSTEM USING IOT, SENSORS, AND FULL-STACK TECHNOLOGIES"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(11)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    field_char_1 = OxmlElement("w:fldChar")
    field_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    field_char_2 = OxmlElement("w:fldChar")
    field_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(field_char_1)
    run._r.append(instr_text)
    run._r.append(field_char_2)


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_separate])
    fallback = paragraph.add_run("Open this file in Microsoft Word and update the Table of Contents field.")
    fallback.italic = True
    fallback.font.size = Pt(12)
    run._r.append(field_end)


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.replace("<br>", "\n").strip()


def add_centered(doc: Document, text: str, size: int = 14, bold: bool = False, italic: bool = False, space_after: float = 6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    add_centered(doc, TITLE, size=16, bold=True, space_after=16)
    add_centered(doc, "A PROJECT REPORT", size=13, bold=True, space_after=12)
    add_centered(doc, "Submitted by", size=13, space_after=6)
    add_centered(doc, "[CANDIDATE NAME(S)]", size=15, bold=True, space_after=16)
    add_centered(doc, "in partial fulfilment of the requirements for the award of the degree", size=12, space_after=6)
    add_centered(doc, "of", size=12, space_after=6)
    add_centered(doc, "BACHELOR OF ENGINEERING / BACHELOR OF TECHNOLOGY", size=14, bold=True, space_after=6)
    add_centered(doc, "in", size=12, space_after=6)
    add_centered(doc, "[BRANCH OF STUDY]", size=13, bold=True, space_after=18)
    add_centered(doc, "KGiSL INSTITUTE OF TECHNOLOGY", size=14, bold=True, space_after=6)
    add_centered(doc, "ANNA UNIVERSITY: CHENNAI 600 025", size=14, bold=True, space_after=8)
    add_centered(doc, "[MONTH AND YEAR]", size=13, bold=True)
    doc.add_page_break()


def add_certificate(doc: Document) -> None:
    add_centered(doc, "KGiSL INSTITUTE OF TECHNOLOGY", size=18, bold=True, space_after=20)
    add_centered(doc, "BONAFIDE CERTIFICATE", size=16, bold=True, space_after=24)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(18)
    p.add_run(
        f'Certified that this project report entitled “{TITLE.title()}” is the bonafide work of '
        '“[CANDIDATE NAME(S)]” who carried out the project work under my supervision.'
    )
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    set_cell_text(table.cell(0, 0), "[Signature of Head of Department]\nSIGNATURE", True)
    set_cell_text(table.cell(0, 1), "[Signature of Supervisor]\nSIGNATURE", True)
    set_cell_text(table.cell(1, 0), "[Name]\nHEAD OF THE DEPARTMENT\n[Department]\n[Full Department and College Address]")
    set_cell_text(table.cell(1, 1), "[Name]\nSUPERVISOR\n[Academic Designation]\n[Department]\n[Full Department and College Address]")
    doc.add_paragraph()
    examiner = doc.add_table(rows=1, cols=2)
    set_cell_text(examiner.cell(0, 0), "INTERNAL EXAMINER", True)
    set_cell_text(examiner.cell(0, 1), "EXTERNAL EXAMINER", True)
    doc.add_page_break()


def add_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = [clean_inline(cell) for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r"[-: ]+", cell) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for index, row_data in enumerate(rows):
        row = table.add_row()
        for column in range(width):
            text = row_data[column] if column < len(row_data) else ""
            set_cell_text(row.cells[column], text, bold=index == 0)
            if index == 0:
                set_cell_shading(row.cells[column], "D9EAD3")
        if index == 0:
            set_repeat_table_header(row)
    doc.add_paragraph()


def add_figure(doc: Document, markdown_line: str) -> None:
    match = re.match(r"!\[([^]]+)\]\(([^)]+)\)", markdown_line.strip())
    if not match:
        return
    caption, relative_path = match.groups()
    image_path = ROOT / relative_path
    if image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(6.2))
        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_p.add_run(clean_inline(caption))
        caption_run.bold = True
        caption_run.font.size = Pt(12)


def add_body_paragraph(doc: Document, text: str, indent: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.35)
    run = p.add_run(clean_inline(text))
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14)


def add_heading(doc: Document, text: str, level: int, insert_page_break: bool = True) -> None:
    if level == 1 and insert_page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(10)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        size = 16
    elif level == 2:
        size = 14
    else:
        size = 14
    run = p.add_run(clean_inline(text).upper() if level == 1 else clean_inline(text))
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_report_body(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(index for index, line in enumerate(lines) if line.strip() == "# ABSTRACT")
    index = start
    first_top_level_heading = True
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip() or line.strip() == "---":
            index += 1
            continue
        if line.startswith("!["):
            add_figure(doc, line)
            index += 1
            continue
        if line.startswith("# "):
            heading = line[2:].strip()
            add_heading(doc, heading, 1, insert_page_break=not first_top_level_heading)
            first_top_level_heading = False
            if heading == "TABLE OF CONTENTS":
                add_toc_field(doc.add_paragraph())
                while index + 1 < len(lines) and (not lines[index + 1].startswith("# ")):
                    index += 1
                continue
            index += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
            index += 1
            continue
        if line.strip().startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(clean_inline(line[2:]))
            run.italic = True
            run.font.size = Pt(13)
            index += 1
            continue
        if line.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.line_spacing = 1
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            index += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            p.add_run(clean_inline(line[2:]))
            index += 1
            continue
        add_body_paragraph(doc, line)
        index += 1


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_certificate(doc)
    add_report_body(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
